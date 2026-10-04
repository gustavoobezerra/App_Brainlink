from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).parent
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.85)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(8)
doc.styles['Normal'].paragraph_format.line_spacing = 1.1
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(16)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
for line in (root / 'Projeto_BrainLink_reuniao.md').read_text(encoding='utf-8').splitlines():
    if not line.strip():
        continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    else:
        p = doc.add_paragraph(line)
        p.paragraph_format.widow_control = True
doc.core_properties.author = 'Gustavo Bezerra'
doc.core_properties.title = 'Projeto BrainLink para reunião com a Psicologia'
doc.core_properties.subject = 'Proposta técnica e colaboração acadêmica no PROBIC'
doc.save(root / 'Projeto_BrainLink_reuniao.docx')
print(root / 'Projeto_BrainLink_reuniao.docx')
